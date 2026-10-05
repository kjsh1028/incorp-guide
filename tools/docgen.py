# -*- coding: utf-8 -*-
"""고객 문서 생성기 v0.3 — 표준 템플릿({{빈칸}})에 사건 값을 채운다. 서식은 템플릿 그대로.
  python tools/docgen.py request  case.json outdir   자료요청서(필요 정보 및 서류 안내)
  python tools/docgen.py signing  case.json outdir   서명 및 공증 서류 안내
  python tools/docgen.py progress case.json outdir   진행 보고
  python tools/docgen.py read     case.json 회신본.docx   자료요청서 회신 판독 -> JSON 출력
채우지 못한 빈칸은 노란색으로 남고 목록으로 알려 준다."""
import sys, json, re, zipfile, io, os, datetime
from xml.sax.saxutils import escape

HERE = os.path.dirname(os.path.abspath(__file__))
WD = {"ko": "월화수목금토일", "zh": "一二三四五六日"}
NAMES = {
 ("request", "ko"): "한국자회사설립_필요정보및서류안내_{short}_KR_{ymd}.docx",
 ("request", "zh"): "韩国子公司设立所需信息及文件清单_{short}_CN_{ymd}.docx",
 ("signing", "ko"): "서명및공증서류안내_{short}_KR_{ymd}.docx",
 ("signing", "zh"): "待签署及公证文件清单_{short}_CN_{ymd}.docx",
 ("progress", "ko"): "진행보고_{short}한국자회사설립_제{no}기_KR_{ymd}.docx",
 ("progress", "zh"): "进度报告_{short}韩国子公司设立_第{no}期_CN_{ymd}.docx",
}
RUN = re.compile(r"<w:r[ >](?:(?!</w:r>).)*</w:r>", re.S)

def L(v, lang):
    """값이 {"ko":..,"zh":..} 이면 해당 언어를, 아니면 그대로."""
    if isinstance(v, dict): return v.get(lang, v.get("ko") or v.get("zh") or "")
    return "" if v is None else v
def xt(s): return escape(str(s)).replace("\n", '</w:t><w:br/><w:t xml:space="preserve">')
def fdate(iso, lang, weekday=True):
    d = datetime.date.fromisoformat(iso)
    if lang == "ko": return "%d. %d. %d." % (d.year, d.month, d.day) + ("(%s)" % WD["ko"][d.weekday()] if weekday else "")
    return "%d年%d月%d日" % (d.year, d.month, d.day) + ("(星期%s)" % WD["zh"][d.weekday()] if weekday else "")
def fdue(iso, lang):
    d = datetime.date.fromisoformat(iso)
    return "%d월 %d일" % (d.month, d.day) if lang == "ko" else "%d月%d日" % (d.month, d.day)

def fill(x, vals):
    def rep(m):
        run = m.group(0)
        for k in re.findall(r"\{\{(\w+)\}\}", run):
            if vals.get(k) not in (None, ""): run = run.replace("{{%s}}" % k, xt(vals[k]))
        if "{{" not in run: run = re.sub(r"<w:highlight[^>]*/>", "", run)
        return run
    return RUN.sub(rep, x)
def bounds(x, i, tag):
    s = max(x.rfind("<w:%s>" % tag, 0, i), x.rfind("<w:%s " % tag, 0, i)); e = x.find("</w:%s>" % tag, i) + len(tag) + 5
    return s, e
def drop(x, needle, tag):
    i = x.find(needle)
    if i < 0: return x
    s, e = bounds(x, i, tag); return x[:s] + x[e:]
def drop_note(x):
    for n in ("발송 전 이 줄 삭제", "发送前删除本行"):
        if n in x: return drop(x, n, "p")
    return x
def repeat(x, key, tag, n_slots, items, render):
    """key_1 이 든 단락/행을 본으로 삼아 items 개수만큼 만들고 나머지 칸은 지운다."""
    i = x.find("{{%s_1}}" % key); s, e = bounds(x, i, tag); proto = x[s:e]
    for k in range(2, n_slots + 1): x = drop(x, "{{%s_%d}}" % (key, k), tag)
    i = x.find("{{%s_1}}" % key); s, e = bounds(x, i, tag)
    out = "".join(render(proto, k, it) for k, it in enumerate(items, 1))
    return x[:s] + out + x[e:]

def common(case, lang):
    f = case["firm"]; sfx = "ko" if lang == "ko" else "cn"
    v = {"client_name_cn": case["client_name_cn"], "receiver_tel": f["tel"], "receiver_email": f["email"],
         "date_" + sfx: fdate(case["date"], lang), "attorney_" + sfx: L(f["attorney"], lang), "receiver_" + sfx: L(f["receiver"], lang)}
    if case.get("reply_by"): v["reply_due_" + sfx] = fdue(case["reply_by"], lang)
    return v

def build(kind, case, lang):
    src = zipfile.ZipFile(os.path.join(HERE, "templates", "%s_%s.docx" % (kind, lang)))
    x = src.read("word/document.xml").decode("utf-8")
    v = common(case, lang); x = drop_note(x)
    if kind == "request":
        opts = [L(o["label"], lang) for o in case.get("investor_options", [])]
        v["investor_option_2"] = "　□ ".join(opts) if opts else None
        if not opts:  # 다른 출자 후보가 없으면 그 선택지 자체를 뺀다
            x = re.sub(r"<w:r[ >](?:(?!</w:r>).)*?\{\{investor_option_2\}\}(?:(?!</w:r>).)*?</w:r>", "", x, flags=re.S)
            x = x.replace("　□ </w:t>", "</w:t>", 1) if "　□ </w:t>" in x else x
    if kind == "signing":
        off = case.get("officers", {})
        if not off.get("other_directors"): x = drop(x, ">A5<", "tr")
        if not off.get("auditor"): x = drop(x, ">A6<", "tr")
    if kind == "progress":
        r = case["report"]; v["report_no"] = r["no"]; v["next_plan"] = L(r.get("next_plan"), lang)
        none = "(없음)" if lang == "ko" else "（无）"
        done = [L(d, lang) for d in r.get("done", [])] or [none]
        x = repeat(x, "done", "p", 2, done, lambda p, k, it: p.replace(">1.<", ">%d.<" % k, 1).replace("{{done_1}}", "{{done_%d}}" % k))
        for k, it in enumerate(done, 1): v["done_%d" % k] = it
        reqs = r.get("requests", []) or [{"item": none, "note": ""}]
        def row(p, k, it):
            p = re.sub(r"(<w:t[^>]*>)1(</w:t>)", r"\g<1>%d\2" % k, p, count=1)
            return p.replace("{{request_1}}", "{{request_%d}}" % k).replace("{{note_1}}", "{{note_%d}}" % k)
        x = repeat(x, "request", "tr", 3, reqs, row)
        for k, it in enumerate(reqs, 1): v["request_%d" % k] = L(it["item"], lang); v["note_%d" % k] = L(it.get("note"), lang) or " "
        for k, st in enumerate(r.get("schedule", []), 1):
            v["t%d" % k] = L(st.get("when"), lang) or " "
            idx = {"done": 0, "doing": 1, "todo": 2}.get(st.get("status"))
            if idx is not None:
                i = x.find("{{t%d}}" % k); s, e = bounds(x, i, "tr"); rowx = x[s:e]
                parts = rowx.split("□")
                if len(parts) == 4: rowx = parts[0] + "".join(("■" if j == idx else "□") + parts[j + 1] for j in range(3))
                x = x[:s] + rowx + x[e:]
    x = fill(x, v)
    left = sorted(set(re.findall(r"\{\{(\w+)\}\}", re.sub(r"<[^>]+>", "", x))))
    buf = io.BytesIO(); dst = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED)
    for it in src.infolist(): dst.writestr(it, x.encode("utf-8") if it.filename == "word/document.xml" else src.read(it.filename))
    dst.close(); return buf.getvalue(), left

def make(kind, case, outdir):
    os.makedirs(outdir, exist_ok=True); res = []
    for lang in case.get("languages", ["zh", "ko"]):
        data, left = build(kind, case, lang)
        name = NAMES[(kind, lang)].format(short=case["client_short"], ymd=case["date"].replace("-", ""), no=case.get("report", {}).get("no", ""))
        p = os.path.join(outdir, name); open(p, "wb").write(data); res.append((p, left))
    return res

# ---------------- 자료요청서 회신 판독 ----------------
UNCHK = "□☐"; CHK = "☑☒■√✓✔●▣◼"
def ctext(c): return "\n".join(p.text for p in c.paragraphs).strip()
def options(t): return [(m.group(1) in CHK, m.group(2).strip("　 ")) for m in re.finditer("([%s%s])\\s*([^%s%s\n]*)" % (UNCHK, CHK, UNCHK, CHK), t)]
def blank(s): return re.sub(r"[_＿\s　]", "", s or "") == ""
def money(s):
    s = s.replace(",", "").replace(" ", "")
    m = re.search(r"(\d+(?:\.\d+)?)(亿|억)", s)
    if m:
        v = float(m.group(1)) * 1e8; m2 = re.search(r"(?:亿|억)(\d+(?:\.\d+)?)(万|만)", s)
        return int(v + (float(m2.group(1)) * 1e4 if m2 else 0))
    m = re.search(r"(\d+(?:\.\d+)?)(万|만)", s)
    if m: return int(float(m.group(1)) * 1e4)
    m = re.search(r"(\d{6,})", s); return int(m.group(1)) if m else None

def read(case, path):
    from docx import Document
    T = Document(path).tables; got, review = {}, []
    def pick(key, text):
        o = options(text); sel = [i for i, (c, _) in enumerate(o) if c]
        if not sel: review.append(key + ": 체크 표시를 읽지 못함"); return None, o
        if len(sel) > 1: review.append(key + ": 둘 이상 체크됨")
        return sel[0], o
    info = {}
    for r in T[0].rows[1:]:
        m = re.search(r"(\d+)\.", ctext(r.cells[0]))
        if m: info[int(m.group(1))] = ctext(r.cells[1])
    so = [{"label": case["client_name_cn"], "type": case.get("client_shareholder_type", "cn")}] + case.get("investor_options", [])
    i, o = pick("1 출자주체", info.get(1, ""))
    if i is not None:
        if i < len(so): got["shareholder"] = L(so[i]["label"], "ko"); got["shareholder_type"] = so[i]["type"]
        else: got["shareholder"] = "기타: " + o[i][1]; review.append("1 출자주체: '기타' 선택. 유형을 직접 판단")
    t = info.get(2, "")
    m = re.search(r"(?:후보|备选)\s*1\s*(.*?)[　\s]*(?:후보|备选)\s*2\s*(.*?)\s*(?:\n|영문|英文)", t, re.S)
    if m: got["name_kr_1"], got["name_kr_2"] = [None if blank(g) else g.strip(" _　") for g in m.groups()]
    m = re.search(r"(?:영문|英文)\s*[:：]\s*(.*)", t)
    if m and not blank(m.group(1)): got["name_en"] = m.group(1).strip(" _")
    if not got.get("name_kr_1"): review.append("2 상호: 한글 상호 미기입")
    t3 = re.split(r"[(（]", info.get(3, ""))[0]; got["capital_krw"] = money(t3)
    if got["capital_krw"] is None: review.append("3 자본금: 미기입 또는 판독 불가")
    i, o = pick("4 1주의 금액", info.get(4, ""))
    if i is not None: got["par_value"] = [5000, 10000][i] if i < 2 else o[i][1]
    bp = info.get(5, "").split("※")[0].strip(); got["business_purposes"] = bp or None
    if not bp: review.append("5 사업목적: 미기입")
    ln = info.get(6, "").split("\n") + ["", ""]
    for key, line, vals, name in (("6 본점 지역", ln[0], ["seoul", "overcrowding_other", "outside", None], "head_office_region"),
                                  ("6 사무실 형태", ln[1], ["independent_lease", "shared_private_room", "kotra_ikp", None], "office_type")):
        i, o = pick(key, line)
        if i is not None and i < len(vals): got[name] = vals[i]
    i, o = pick("7 공고방법", info.get(7, "")); got["notice_method"] = None if i is None else ["homepage_then_newspaper", "newspaper"][min(i, 1)]
    i, o = pick("8 ODI", info.get(8, "")); got["odi_status"] = None if i is None else ["done", "in_progress", "not_started"][min(i, 2)]
    i, o = pick("9 선호 은행", info.get(9, ""))
    if i is not None: got["preferred_bank"] = (re.split(r"[:：]", o[i][1])[-1].strip(" _") or None) if i == 0 else None
    c10 = re.sub(r"^.*?[:：]", "", info.get(10, ""), count=1, flags=re.S).strip(); got["contact"] = c10 or None
    if not c10: review.append("10 연락 담당자: 미기입")
    off = {}; keys = ["name_en", "name_zh", "nationality_dob", "passport_no", "address", "resident_in_korea"]
    for col, role in ((1, "rep_director"), (2, "director"), (3, "auditor")):
        d = {}
        for k, r in zip(keys, T[1].rows[1:]):
            val = ctext(r.cells[col])
            if k == "resident_in_korea":
                sel = [j for j, (c, _) in enumerate(options(val)) if c]; d[k] = None if not sel else sel[0] == 0
            else: d[k] = val or None
        if any(d[k] for k in keys[:5]): off[role] = d
    got["officers"] = off
    if "rep_director" not in off: review.append("임원: 대표이사 정보 미기입")
    elif not off["rep_director"].get("address"): review.append("임원: 대표이사 주소 미기입(★ 항목)")
    for t in T:
        if len(t.rows) == 1 and len(t.rows[0].cells) == 1:
            o = options(ctext(t.rows[0].cells[0]))
            if len(o) == 4:
                sel = [j for j, (c, _) in enumerate(o) if c]
                got["pre_registration_cost_method"] = None if not sel else ["yulchon_advance", "client_remit", "korean_affiliate", "other"][sel[0]]
                if not sel: review.append("등기 전 비용 지급 방식: 미선택")
    inp = {"dir": 1 + (1 if "director" in off else 0), "aud": "auditor" in off}
    if got.get("shareholder_type"): inp["sh"] = got["shareholder_type"]
    if got.get("capital_krw"): inp["cap"] = got["capital_krw"]
    if got.get("head_office_region"): inp["zone"] = "other" if got["head_office_region"] == "outside" else "over"
    if got.get("office_type"): inp["off"] = "indep"
    if got.get("office_type") and got["office_type"] != "independent_lease":
        review.append("6 사무실 형태: 회신은 '%s'인데 화면에는 '독립 사무실'로 넣었음. 직접 확인" % {"shared_private_room": "공유오피스 독립실", "kotra_ikp": "KOTRA IKP"}.get(got["office_type"], got["office_type"]))
    rk = (off.get("rep_director") or {}).get("resident_in_korea")
    if rk is not None: inp["d8"] = bool(rk)
    review.append("대표이사 거주지·신분증 유형, 이사가 2명 이상인 경우의 이사 수, 제한 업종 여부는 서식만으로 판단할 수 없어 직접 확인")
    out = dict(case); out.update(answers=got, program_inputs=inp, review=review); return out

if __name__ == "__main__":
    kind = sys.argv[1]; case = json.load(open(sys.argv[2], encoding="utf-8"))
    fp = os.path.join(HERE, "firm_profile.json")
    if "firm" not in case and os.path.exists(fp): case["firm"] = json.load(open(fp, encoding="utf-8"))
    if kind == "read": print(json.dumps(read(case, sys.argv[3]), ensure_ascii=False, indent=1))
    else:
        for p, left in make(kind, case, sys.argv[3]): print(p, ("| 남은 빈칸: " + ", ".join(left)) if left else "")
