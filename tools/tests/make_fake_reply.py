# 가상 회신본 만들기 (시험용, 실제 고객 정보 아님)
#   python tools/tests/make_fake_reply.py <docgen request 출력 폴더> <저장할 파일>
import docx,glob,sys
d=docx.Document(glob.glob(sys.argv[1]+"/*_KR_*.docx")[0])
def nth_replace(par, old, new, n=1):
    """단락 안에서 n번째 old 를 new 로 (run 경계 안에서)."""
    k=0
    for r in par.runs:
        t=r.text; i=-1
        while True:
            i=t.find(old,i+1)
            if i<0: break
            k+=1
            if k==n: r.text=t[:i]+new+t[i+len(old):]; return True
    return False
def cellp(t,ri,ci,pi=0): return d.tables[t].rows[ri].cells[ci].paragraphs[pi]
T0=0
nth_replace(cellp(T0,1,1),"□","☑",1)                       # 출자주체: 중국 법인
p=d.tables[T0].rows[2].cells[1].paragraphs
nth_replace(p[0],"________","한빛테크",1); nth_replace(p[0],"________","한빛",1)
for q in p:
    if "영문" in q.text: nth_replace(q,"________________","Hanbit Tech Co., Ltd.",1)
nth_replace(cellp(T0,3,1),"________","300,000,000",1)       # 자본금
nth_replace(cellp(T0,4,1),"□","☑",2)                        # 1주 10,000원
c5=d.tables[T0].rows[5].cells[1].paragraphs; c5[0].add_run("1. 소프트웨어 개발 및 공급"); c5[1].add_run("2. 위 각호에 부대하는 사업")
c6=d.tables[T0].rows[6].cells[1].paragraphs; nth_replace(c6[0],"□","☑",1); nth_replace(c6[1],"□","☑",2)  # 서울, 공유오피스 독립실
c7=d.tables[T0].rows[7].cells[1].paragraphs; nth_replace(c7[-1],"□","☑",1)   # 신문
nth_replace(cellp(T0,8,1),"□","☑",2)                        # ODI 진행 중
nth_replace(cellp(T0,9,1),"□","☑",1); nth_replace(cellp(T0,9,1),"________","하나은행",1)
cellp(T0,10,1).add_run(" 가상 담당자, 법무팀장, +86 000 0000 0000, test@example.com")
vals={1:["TEST REP","测试代表"],2:["TEST DIR","测试董事"]}
for col,(en,zh) in vals.items():
    rows=d.tables[1].rows
    for ri,v in zip([1,2,3,4,5],[en,zh,"중국 / 1980-01-01","X0000000"+str(col),"中国上海市测试路1号"]): rows[ri].cells[col].paragraphs[0].add_run(v)
    nth_replace(rows[6].cells[col].paragraphs[0],"□","☑",2 if col==1 else 1)   # 대표 아니오, 이사 예
for q in d.tables[6].rows[0].cells[0].paragraphs:
    if "미리 송금" in q.text: nth_replace(q,"□","☑",1)
d.save(sys.argv[2])
